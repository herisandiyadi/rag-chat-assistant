"""Document service for upload, management, and ingestion."""

import logging
import uuid
from typing import List, Dict, Any, Optional
from io import BytesIO

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_

from src.core.db import get_db
from src.models.document import Document
from src.models.user import User
from src.rag.vector_search import get_vector_search
from src.schemas import DocumentCreate, DocumentOut

logger = logging.getLogger(__name__)


class DocumentService:
    """Service for document management."""
    
    @staticmethod
    async def list_documents(
        db: AsyncSession,
        user_id: str,
        skip: int = 0,
        limit: int = 100,
    ) -> Dict[str, Any]:
        """List documents user has access to."""
        user = await db.get(User, user_id)
        if not user:
            raise ValueError("User not found")
        
        # Build query based on user role
        query = select(Document)
        
        if user.role_type == "dept_admin":
            query = query.where(Document.department_id == user.department_id)
        elif user.role_type == "user":
            # Regular users see only their department's documents
            query = query.where(Document.department_id == user.department_id)
        
        # Get total count
        count_query = select(Document)
        if user.role_type == "dept_admin":
            count_query = count_query.where(Document.department_id == user.department_id)
        elif user.role_type == "user":
            count_query = count_query.where(Document.department_id == user.department_id)
        
        count_result = await db.execute(count_query)
        total = len(count_result.scalars().all())
        
        query = query.offset(skip).limit(limit)
        result = await db.execute(query)
        documents = result.scalars().all()
        
        return {
            "total": total,
            "documents": [
                {
                    "id": doc.id,
                    "doc_id": doc.doc_id,
                    "judul": doc.judul,
                    "department_id": doc.department_id,
                    "min_level": doc.min_level,
                    "hidden_existence": doc.hidden_existence,
                    "versi": doc.versi,
                    "status": doc.status,
                    "created_at": doc.created_at,
                }
                for doc in documents
            ],
        }
    
    @staticmethod
    async def upload_document(
        db: AsyncSession,
        user_id: str,
        file,
        department: Optional[str] = None,
        min_level: int = 1,
        hidden_existence: bool = False,
    ) -> Dict[str, Any]:
        """Upload a new document with metadata and trigger ingestion."""
        user = await db.get(User, user_id)
        if not user:
            raise ValueError("User not found")
        
        # Validate user has permission
        if user.role_type not in ["super_admin", "dept_admin"]:
            raise PermissionError("Not authorized to upload documents")
        
        # Generate doc_id
        doc_id = str(uuid.uuid4())
        
        # Extract filename
        filename = getattr(file, "filename", "unknown")
        
        # Create document record
        document = Document(
            doc_id=doc_id,
            judul=filename.rsplit(".", 1)[0],
            department_id=user.department_id if user.role_type == "dept_admin" else 1,
            min_level=min_level,
            hidden_existence=hidden_existence,
            owner_admin_id=user.id,
            versi=1,
            nama_file=filename,
        )
        
        db.add(document)
        await db.commit()
        await db.refresh(document)
        
        # Read file content
        file_content = await file.read()
        
        # Extract text (simplified - in production, use proper document parsing)
        text = DocumentService._extract_text_from_file(file_content, filename)
        
        # Chunk text
        chunks = DocumentService._chunk_text(text, chunk_size=500)
        
        # Ingest to vector database
        vector_search = get_vector_search()
        metadata = {
            "department_id": str(document.department_id),
            "doc_id": doc_id,
            "min_level": min_level,
            "hidden_existence": hidden_existence,
            "doc_title": filename,
        }
        vector_search.add_chunks(doc_id, chunks, metadata)
        
        return {
            "id": document.id,
            "doc_id": doc_id,
            "judul": document.judul,
            "chunk_count": len(chunks),
            "min_level": min_level,
            "hidden_existence": hidden_existence,
        }
    
    @staticmethod
    async def update_document(
        db: AsyncSession,
        user_id: str,
        doc_id: int,
        min_level: Optional[int] = None,
        hidden_existence: Optional[bool] = None,
    ) -> Optional[Document]:
        """Update document metadata without re-embedding."""
        user = await db.get(User, user_id)
        if not user:
            raise ValueError("User not found")
        
        # Check permission
        if user.role_type not in ["super_admin", "dept_admin"]:
            raise PermissionError("Not authorized to update documents")
        
        # Get document
        result = await db.execute(select(Document).where(Document.id == doc_id))
        document = result.scalar_one_or_none()
        if not document:
            return None
        
        # Update fields
        if min_level is not None:
            document.min_level = min_level
        if hidden_existence is not None:
            document.hidden_existence = hidden_existence
        
        await db.commit()
        await db.refresh(document)
        
        return document
    
    @staticmethod
    async def delete_document(
        db: AsyncSession,
        user_id: str,
        doc_id: int,
    ) -> bool:
        """Delete a document."""
        user = await db.get(User, user_id)
        if not user:
            raise ValueError("User not found")
        
        # Check permission
        if user.role_type not in ["super_admin", "dept_admin"]:
            raise PermissionError("Not authorized to delete documents")
        
        # Get document
        result = await db.execute(select(Document).where(Document.id == doc_id))
        document = result.scalar_one_or_none()
        if not document:
            return False
        
        await db.delete(document)
        await db.commit()
        return True
    
    @staticmethod
    def _extract_text_from_file(content: bytes, filename: str) -> str:
        """Extract text from different file types."""
        try:
            if filename.endswith(".pdf"):
                import PyPDF2
                pdf_reader = PyPDF2.PdfReader(BytesIO(content))
                return "\n".join([
                    page.extract_text() or ""
                    for page in pdf_reader.pages
                ])
            elif filename.endswith(".docx"):
                from docx import Document as DocxDocument
                doc = DocxDocument(BytesIO(content))
                return "\n".join([
                    paragraph.text
                    for paragraph in doc.paragraphs
                ])
            else:
                # Default: treat as plain text
                return content.decode("utf-8")
        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {e}")
            return ""
    
    @staticmethod
    def _chunk_text(text: str, chunk_size: int = 500) -> List[str]:
        """Split text into chunks."""
        words = text.split()
        chunks = []
        current_chunk = []
        current_length = 0
        
        for word in words:
            word_length = len(word) + 1  # +1 for space
            if current_length + word_length > chunk_size and current_chunk:
                chunks.append(" ".join(current_chunk))
                current_chunk = [word]
                current_length = word_length
            else:
                current_chunk.append(word)
                current_length += word_length
        
        if current_chunk:
            chunks.append(" ".join(current_chunk))
        
        return chunks


def get_document_service() -> DocumentService:
    """Get document service instance."""
    return DocumentService()
